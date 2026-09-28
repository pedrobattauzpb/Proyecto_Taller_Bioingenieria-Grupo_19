"""
Generador de archivos Excel (.xlsx) estructurados para inspecciones técnicas
y reportes ejecutivos institucionales utilizando openpyxl.
"""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.reports import InspectionReportData, ExecutiveReportData


# Estilos de encabezado y celdas
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1D5652")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="16212B")
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
REGULAR_FONT = Font(name="Calibri", size=10)
SMALL_FONT = Font(name="Calibri", size=9, color="55636E")

PRIMARY_FILL = PatternFill(start_color="2F6F6B", end_color="2F6F6B", fill_type="solid")
PRIMARY_SOFT_FILL = PatternFill(start_color="E3EFEE", end_color="E3EFEE", fill_type="solid")
OK_FILL = PatternFill(start_color="E5F4EC", end_color="E5F4EC", fill_type="solid")
WARN_FILL = PatternFill(start_color="FBF0DD", end_color="FBF0DD", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
GRAY_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="D1D9E0"),
    right=Side(style="thin", color="D1D9E0"),
    top=Side(style="thin", color="D1D9E0"),
    bottom=Side(style="thin", color="D1D9E0"),
)


def _auto_fit_columns(ws, min_width=12, max_width=50):
    """Ajusta automáticamente el ancho de las columnas según su contenido."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len and "\n" not in val:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


class InspectionExcelGenerator:
    """Generador de hojas de cálculo Excel (.xlsx) para inspecciones técnicas."""

    def generate(self, data: InspectionReportData) -> BytesIO:
        wb = Workbook()

        # Hoja 1: Resumen General
        ws_resumen = wb.active
        ws_resumen.title = "Resumen de Inspección"
        ws_resumen.views.sheetView[0].showGridLines = True

        ws_resumen["A1"] = "ACTA DE INSPECCIÓN TÉCNICA — GASES MEDICINALES"
        ws_resumen["A1"].font = TITLE_FONT
        ws_resumen["A2"] = f"Inspección #{data.inspection_id} · Protocolo: {data.template_title} · Estado: {data.status}"
        ws_resumen["A2"].font = SMALL_FONT

        # Datos Generales
        meta = [
            ("Hospital:", data.hospital_name, "Inspector:", data.inspector_name),
            ("Código Hospital:", data.hospital_code, "Fecha Inicio:", data.started_at.strftime("%d/%m/%Y %H:%M")),
            ("Sector / Piso:", f"{data.sector_name} (Piso {data.sector_floor or 'N/A'})", "Fecha Cierre:", data.completed_at.strftime("%d/%m/%Y %H:%M") if data.completed_at else "En Curso"),
            ("Activo:", data.asset_name, "TAG Activo:", data.asset_tag),
            ("Tipo Componente:", data.asset_type, "N° Serie:", data.asset_serial or "No informado"),
            ("Plantilla Normativa:", f"{data.template_title} v{data.template_version}", "QR Activo:", data.asset_qr or "No asignado"),
        ]

        row = 4
        for label1, val1, label2, val2 in meta:
            ws_resumen.cell(row=row, column=1, value=label1).font = BOLD_FONT
            ws_resumen.cell(row=row, column=2, value=val1).font = REGULAR_FONT
            ws_resumen.cell(row=row, column=3, value=label2).font = BOLD_FONT
            ws_resumen.cell(row=row, column=4, value=val2).font = REGULAR_FONT
            for col in range(1, 5):
                ws_resumen.cell(row=row, column=col).border = THIN_BORDER
            row += 1

        # Resumen Compliance
        row += 2
        ws_resumen.cell(row=row, column=1, value="DICTAMEN DE CONFORMIDAD NORMATIVA").font = SECTION_FONT
        row += 1

        cs = data.compliance_summary
        comp_metrics = [
            ("Porcentaje de Cumplimiento:", f"{cs.compliance_percentage:.1f}%"),
            ("Estado General:", "PLENAMENTE CONFORME" if cs.is_fully_compliant else "NO CONFORME (REQUIERE ACCIÓN)"),
            ("Total de Ítems Evaluados:", f"{cs.evaluated_items} de {cs.total_items}"),
            ("Ítems Conformes:", cs.compliant_items),
            ("No Conformidades Totales:", cs.non_compliant_items),
            ("Desvíos Críticos:", cs.critical_deviations_count),
            ("Desvíos Mayores:", cs.major_deviations_count),
            ("Advertencias (Zona 10%):", cs.warning_items),
        ]

        for label, val in comp_metrics:
            ws_resumen.cell(row=row, column=1, value=label).font = BOLD_FONT
            c2 = ws_resumen.cell(row=row, column=2, value=val)
            c2.font = BOLD_FONT if "Porcentaje" in label or "Estado" in label else REGULAR_FONT
            ws_resumen.cell(row=row, column=1).border = THIN_BORDER
            c2.border = THIN_BORDER
            row += 1

        if data.notes:
            row += 2
            ws_resumen.cell(row=row, column=1, value="Observaciones Generales del Inspector:").font = BOLD_FONT
            ws_resumen.cell(row=row + 1, column=1, value=data.notes).font = REGULAR_FONT

        _auto_fit_columns(ws_resumen)

        # Hoja 2: Detalle de Ítems
        ws_items = wb.create_sheet(title="Checklist Detallado")
        ws_items.views.sheetView[0].showGridLines = True

        headers = [
            "Código",
            "Punto de Control",
            "Referencia Normativa",
            "Tipo",
            "Obligatorio",
            "Valor Medido",
            "Unidad",
            "Esperado / Rango",
            "Dictamen",
            "Severidad",
            "Detalle del Desvío",
            "Observaciones del Auditor",
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws_items.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = PRIMARY_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

        for row_idx, it in enumerate(data.items, 2):
            if it.input_type == "BOOLEAN":
                val = "Sí" if it.val_boolean else ("No" if it.val_boolean is False else "Sin evaluar")
            elif it.input_type == "NUMERIC":
                val = str(it.val_numeric) if it.val_numeric is not None else "Sin evaluar"
            else:
                val = it.val_text or "Sin texto"

            # Rango esperado
            if it.input_type == "NUMERIC":
                if it.min_value is not None and it.max_value is not None:
                    exp = f"{it.min_value} - {it.max_value}"
                elif it.min_value is not None:
                    exp = f">= {it.min_value}"
                elif it.max_value is not None:
                    exp = f"<= {it.max_value}"
                else:
                    exp = "Reglamentario"
            elif it.input_type == "BOOLEAN":
                exp = "Sí (PASA)"
            else:
                exp = "Dato requerido"

            row_data = [
                it.code,
                it.title,
                it.normative_ref,
                it.input_type,
                "Sí" if it.is_mandatory else "No",
                val,
                it.unit or "",
                exp,
                it.compliance_status,
                it.severity,
                it.deviation_detail or "",
                it.observations or "",
            ]

            fill = GRAY_FILL if row_idx % 2 == 0 else PatternFill(fill_type=None)
            if it.compliance_status == "NON_COMPLIANT":
                fill = CRITICAL_FILL if it.severity == "CRITICAL" else WARN_FILL
            elif it.compliance_status == "COMPLIANT":
                fill = OK_FILL

            for col_idx, val_cell in enumerate(row_data, 1):
                c = ws_items.cell(row=row_idx, column=col_idx, value=val_cell)
                c.font = REGULAR_FONT
                c.border = THIN_BORDER
                if fill.fill_type:
                    c.fill = fill

        _auto_fit_columns(ws_items)

        # Hoja 3: Evidencias
        if data.evidences:
            ws_evid = wb.create_sheet(title="Evidencias Multimedia")
            ws_evid.views.sheetView[0].showGridLines = True

            evid_headers = ["ID", "Tipo Archivo", "Identificador Almacenamiento", "Tamaño (Bytes)", "Fecha Carga", "Cargado Por"]
            for col_idx, h in enumerate(evid_headers, 1):
                cell = ws_evid.cell(row=1, column=col_idx, value=h)
                cell.font = HEADER_FONT
                cell.fill = PRIMARY_FILL
                cell.alignment = Alignment(horizontal="center")
                cell.border = THIN_BORDER

            for row_idx, e in enumerate(data.evidences, 2):
                e_row = [
                    e.id,
                    e.file_type,
                    e.storage_url.split("/")[-1],
                    e.file_size_bytes or 0,
                    e.uploaded_at.strftime("%d/%m/%Y %H:%M:%S"),
                    e.uploaded_by or "Inspector",
                ]
                for col_idx, val_cell in enumerate(e_row, 1):
                    c = ws_evid.cell(row=row_idx, column=col_idx, value=val_cell)
                    c.font = REGULAR_FONT
                    c.border = THIN_BORDER

            _auto_fit_columns(ws_evid)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer


class ExecutiveExcelGenerator:
    """Generador de informes ejecutivos institucionales en Excel."""

    def generate(self, data: ExecutiveReportData) -> BytesIO:
        wb = Workbook()

        # Hoja 1: Indicadores Globales
        ws_kpi = wb.active
        ws_kpi.title = "KPIs Globales"
        ws_kpi.views.sheetView[0].showGridLines = True

        ws_kpi["A1"] = f"INFORME EJECUTIVO DE CUMPLIMIENTO — {data.hospital_name}"
        ws_kpi["A1"].font = TITLE_FONT
        ws_kpi["A2"] = f"Código: {data.hospital_code} · Generado: {data.generated_at.strftime('%d/%m/%Y %H:%M')}"
        ws_kpi["A2"].font = SMALL_FONT

        kpi_data = [
            ("Sectores Monitoreados:", data.total_sectors),
            ("Total de Activos Registrados:", data.total_assets),
            ("Activos Clínicos Operativos:", data.active_assets),
            ("Inspecciones Totales Realizadas:", data.total_inspections),
            ("Inspecciones Cerradas y Selladas:", data.completed_inspections),
            ("Inspecciones en Curso:", data.in_progress_inspections),
            ("Porcentaje de Cumplimiento Global:", f"{data.global_compliance_percentage:.1f}%"),
        ]

        row = 4
        for label, val in kpi_data:
            ws_kpi.cell(row=row, column=1, value=label).font = BOLD_FONT
            ws_kpi.cell(row=row, column=2, value=val).font = BOLD_FONT if "Cumplimiento" in label else REGULAR_FONT
            ws_kpi.cell(row=row, column=1).border = THIN_BORDER
            ws_kpi.cell(row=row, column=2).border = THIN_BORDER
            row += 1

        _auto_fit_columns(ws_kpi)

        # Hoja 2: Cumplimiento por Sector
        ws_sec = wb.create_sheet(title="Por Sector")
        ws_sec.views.sheetView[0].showGridLines = True

        sec_headers = ["Sector", "Nivel / Piso", "Total Activos", "Inspecciones Totales", "Cerradas", "% Cumplimiento"]
        for col_idx, h in enumerate(sec_headers, 1):
            cell = ws_sec.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = PRIMARY_FILL
            cell.border = THIN_BORDER

        for row_idx, s in enumerate(data.sector_stats, 2):
            s_row = [
                s["sector_name"],
                s["floor_level"],
                s["total_assets"],
                s["total_inspections"],
                s["completed_inspections"],
                f"{s['compliance_percentage']:.1f}%",
            ]
            for col_idx, val_cell in enumerate(s_row, 1):
                c = ws_sec.cell(row=row_idx, column=col_idx, value=val_cell)
                c.font = REGULAR_FONT
                c.border = THIN_BORDER

        _auto_fit_columns(ws_sec)

        # Hoja 3: Cumplimiento por Tipo de Activo
        ws_type = wb.create_sheet(title="Por Tipo de Activo")
        ws_type.views.sheetView[0].showGridLines = True

        type_headers = ["Familia de Activo", "Total Instalado", "Auditorías Realizadas", "% Cumplimiento"]
        for col_idx, h in enumerate(type_headers, 1):
            cell = ws_type.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = PRIMARY_FILL
            cell.border = THIN_BORDER

        for row_idx, t in enumerate(data.asset_type_stats, 2):
            t_row = [
                t["asset_type"],
                t["total_assets"],
                t["total_inspections"],
                f"{t['compliance_percentage']:.1f}%",
            ]
            for col_idx, val_cell in enumerate(t_row, 1):
                c = ws_type.cell(row=row_idx, column=col_idx, value=val_cell)
                c.font = REGULAR_FONT
                c.border = THIN_BORDER

        _auto_fit_columns(ws_type)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
